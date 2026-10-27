from __future__ import annotations

import csv
import hashlib
import hmac
import io
import json
import logging
import os
import re
import secrets
import smtplib
import uuid
from datetime import datetime, timezone
from email.message import EmailMessage
from pathlib import Path
from typing import Any, Iterator
from urllib.parse import quote

import httpx
from docx import Document
from fastapi import Depends, FastAPI, File, Header, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from pydantic import BaseModel, Field
from pypdf import PdfReader
from sqlalchemy import Boolean, DateTime, Float, ForeignKey, Integer, JSON, String, Text, UniqueConstraint, create_engine, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column, relationship, sessionmaker


logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO"))
log = logging.getLogger("neural-ant")

DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///./data/neural_ant.db")
if DATABASE_URL.startswith("postgres://"):
    DATABASE_URL = "postgresql+psycopg://" + DATABASE_URL.removeprefix("postgres://")
elif DATABASE_URL.startswith("postgresql://"):
    DATABASE_URL = "postgresql+psycopg://" + DATABASE_URL.removeprefix("postgresql://")
UPLOAD_DIR = Path(os.getenv("UPLOAD_DIR", "./data/resumes"))
GCS_BUCKET = os.getenv("GCS_BUCKET", "").strip()
APP_ENV = os.getenv("APP_ENV", "development").lower()
SESSION_SECRET = os.getenv("SESSION_SECRET", "")
COOKIE_NAME = "neural_ant_session"
SESSION_MAX_AGE = 60 * 60 * 24 * 14
MAX_RESUME_BYTES = 10 * 1024 * 1024
DIST_DIR = Path(__file__).resolve().parent.parent / "dist"

connect_args = {"check_same_thread": False} if DATABASE_URL.startswith("sqlite") else {}
if DATABASE_URL.startswith("sqlite") and DATABASE_URL != "sqlite:///:memory:":
    Path(DATABASE_URL.removeprefix("sqlite:///" )).parent.mkdir(parents=True, exist_ok=True)
engine = create_engine(DATABASE_URL, connect_args=connect_args, pool_pre_ping=True)
SessionLocal = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
if APP_ENV == "production" and (len(SESSION_SECRET) < 32 or not os.getenv("NEURAL_ANT_PASSWORD")):
    raise RuntimeError("Production requires SESSION_SECRET (32+ characters) and NEURAL_ANT_PASSWORD.")
serializer = URLSafeTimedSerializer(SESSION_SECRET or "local-development-only-change-this", salt="neural-ant-session")


class Base(DeclarativeBase):
    pass


class ProfileRow(Base):
    __tablename__ = "profile"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, default=1)
    data: Mapped[dict[str, Any]] = mapped_column(JSON, default=dict)
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class ResumeRow(Base):
    __tablename__ = "resume"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, default=1)
    filename: Mapped[str] = mapped_column(String(255))
    file_key: Mapped[str] = mapped_column(String(1024))
    file_sha256: Mapped[str] = mapped_column(String(64))
    extracted_text: Mapped[str] = mapped_column(Text)
    uploaded_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class SourceRow(Base):
    __tablename__ = "job_sources"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    ats_type: Mapped[str] = mapped_column(String(24), index=True)
    slug: Mapped[str] = mapped_column(String(160))
    company_name: Mapped[str] = mapped_column(String(255))
    enabled: Mapped[bool] = mapped_column(Boolean, default=True)
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    __table_args__ = (UniqueConstraint("ats_type", "slug", name="uq_source_type_slug"),)


class SponsorRow(Base):
    __tablename__ = "sponsor_evidence"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    company_key: Mapped[str] = mapped_column(String(255), index=True)
    company_name: Mapped[str] = mapped_column(String(255))
    evidence_type: Mapped[str] = mapped_column(String(80), default="public filing")
    evidence_url: Mapped[str] = mapped_column(String(2000))
    evidence_year: Mapped[str] = mapped_column(String(16), default="")
    filing_count: Mapped[int] = mapped_column(Integer, default=0)
    notes: Mapped[str] = mapped_column(Text, default="")
    imported_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class JobRow(Base):
    __tablename__ = "jobs"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    url: Mapped[str] = mapped_column(String(2000), unique=True, index=True)
    company_name: Mapped[str] = mapped_column(String(255), index=True)
    title: Mapped[str] = mapped_column(String(500), index=True)
    source: Mapped[str] = mapped_column(String(32))
    location: Mapped[str] = mapped_column(String(500), default="")
    work_mode: Mapped[str] = mapped_column(String(80), default="Unspecified")
    salary_text: Mapped[str] = mapped_column(String(255), default="Not stated")
    salary_min: Mapped[int | None] = mapped_column(Integer, nullable=True)
    salary_max: Mapped[int | None] = mapped_column(Integer, nullable=True)
    description: Mapped[str] = mapped_column(Text, default="")
    posted_at: Mapped[str] = mapped_column(String(80), default="")
    fit_score: Mapped[int] = mapped_column(Integer, default=0, index=True)
    fit_reasons: Mapped[list[str]] = mapped_column(JSON, default=list)
    sponsorship_status: Mapped[str] = mapped_column(String(40), default="unclear")
    sponsorship_evidence: Mapped[str] = mapped_column(Text, default="")
    sponsorship_url: Mapped[str] = mapped_column(String(2000), default="")
    first_seen: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    last_seen: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    applications: Mapped[list["ApplicationRow"]] = relationship(back_populates="job")


class ApplicationRow(Base):
    __tablename__ = "applications"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    job_id: Mapped[int] = mapped_column(ForeignKey("jobs.id"), index=True)
    stage: Mapped[str] = mapped_column(String(32), default="Shortlist")
    notes: Mapped[str] = mapped_column(Text, default="")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    updated_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))
    job: Mapped[JobRow] = relationship(back_populates="applications")


class DraftRow(Base):
    __tablename__ = "resume_drafts"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    job_id: Mapped[int] = mapped_column(ForeignKey("jobs.id"), index=True)
    content: Mapped[str] = mapped_column(Text)
    model: Mapped[str] = mapped_column(String(160), default="local-match-draft")
    created_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


class AlertRow(Base):
    __tablename__ = "email_alerts"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    job_url: Mapped[str] = mapped_column(String(2000), unique=True)
    sent_at: Mapped[datetime] = mapped_column(DateTime(timezone=True), default=lambda: datetime.now(timezone.utc))


DEFAULT_PROFILE: dict[str, Any] = {
    "name": "",
    "email": "",
    "role": "",
    "skills": [],
    "education": "",
    "target_titles": [],
    "location_preferences": ["United States"],
    "work_modes": ["Remote", "Hybrid", "On-site"],
    "minimum_salary": None,
    "years_experience": None,
    "work_authorization": "",
    "sponsorship_preferences": "",
    "summary": "",
    "email_alerts": False,
    "match_threshold": 55,
}

SUPPORTED_ATS = {"greenhouse", "lever", "ashby"}
STAGES = {"Shortlist", "Applied", "Interview", "Decision", "Archived"}
SENSITIVE_HEADERS = {"host", "x-forwarded-host"}


def get_db() -> Iterator[Session]:
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def ensure_profile(db: Session) -> ProfileRow:
    row = db.get(ProfileRow, 1)
    if row is None:
        row = ProfileRow(id=1, data=DEFAULT_PROFILE.copy())
        db.add(row)
        db.commit()
        db.refresh(row)
    return row


def now_iso(value: datetime | None) -> str | None:
    return value.isoformat() if value else None


def safe_json(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def company_key(name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", name.lower())


def parse_salary(text: str) -> tuple[int | None, int | None]:
    values: list[int] = []
    for raw, suffix in re.findall(r"\$\s*([\d,]+(?:\.\d+)?)\s*(k|m)?", text or "", re.I):
        try:
            number = float(raw.replace(",", ""))
            if suffix.lower() == "k":
                number *= 1000
            elif suffix.lower() == "m":
                number *= 1_000_000
            elif number < 1000:
                continue
            values.append(int(number))
        except ValueError:
            pass
    return (min(values), max(values)) if values else (None, None)


def extract_resume_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
        if suffix == ".docx":
            document = Document(io.BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            return "\n".join(parts).strip()
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read this resume: {exc.__class__.__name__}") from exc
    raise HTTPException(status_code=415, detail="Upload a PDF or DOCX resume.")


def extract_profile(text: str) -> dict[str, Any]:
    lines = [line.strip(" •\t-") for line in text.splitlines() if line.strip()]
    name = ""
    for line in lines[:8]:
        if 2 <= len(line) <= 72 and not re.search(r"@|https?://|\d{3}|resume|curriculum vitae", line, re.I):
            name = line
            break
    email_match = re.search(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", text, re.I)
    skill_catalog = [
        "Python", "Java", "TypeScript", "JavaScript", "SQL", "AWS", "Azure", "GCP", "Docker", "Kubernetes",
        "Terraform", "FastAPI", "React", "Machine Learning", "LLM", "RAG", "NLP", "PyTorch", "TensorFlow",
        "LangChain", "Bedrock", "OpenAI", "PostgreSQL", "MongoDB", "MLOps", "CI/CD",
    ]
    skills = [skill for skill in skill_catalog if re.search(r"\b" + re.escape(skill) + r"\b", text, re.I)]
    return {"name": name, "email": email_match.group(0) if email_match else "", "skills": skills}


def put_resume_file(file_key: str, filename: str, data: bytes) -> str:
    if GCS_BUCKET:
        from google.cloud import storage

        bucket = storage.Client().bucket(GCS_BUCKET)
        key = f"resumes/{file_key}/{Path(filename).name}"
        bucket.blob(key).upload_from_string(data, content_type="application/octet-stream")
        return f"gs://{GCS_BUCKET}/{key}"
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    destination = UPLOAD_DIR / file_key
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_bytes(data)
    return str(destination)


def delete_resume_file(file_ref: str) -> None:
    if file_ref.startswith("gs://"):
        try:
            from google.cloud import storage

            bucket_name, _, key = file_ref[5:].partition("/")
            if not bucket_name or not key:
                raise ValueError("Invalid GCS object reference")
            storage.Client().bucket(bucket_name).blob(key).delete()
        except Exception:
            log.exception("Unable to delete resume object from cloud storage")
    else:
        try:
            Path(file_ref).unlink(missing_ok=True)
        except OSError:
            log.exception("Unable to delete local resume file")


def authenticated(request: Request) -> bool:
    token = request.cookies.get(COOKIE_NAME, "")
    if not token:
        return False
    try:
        session_data = serializer.loads(token, max_age=SESSION_MAX_AGE)
        return session_data.get("sub") == "owner"
    except (BadSignature, SignatureExpired):
        return False


def require_auth(request: Request) -> None:
    if not authenticated(request):
        raise HTTPException(status_code=401, detail="Sign in to Neural Ant.")


def public_job_payload(job: JobRow) -> dict[str, Any]:
    return {
        "id": job.id,
        "title": job.title,
        "company": job.company_name,
        "location": job.location,
        "work_mode": job.work_mode,
        "salary": job.salary_text,
        "salary_min": job.salary_min,
        "salary_max": job.salary_max,
        "description": job.description[:12000],
        "url": job.url,
        "source": job.source,
        "posted_at": job.posted_at,
        "fit_score": job.fit_score,
        "fit_reasons": job.fit_reasons or [],
        "sponsorship_status": job.sponsorship_status,
        "sponsorship_evidence": job.sponsorship_evidence,
        "sponsorship_url": job.sponsorship_url,
        "first_seen": now_iso(job.first_seen),
    }


def parse_greenhouse(slug: str, company: str, client: httpx.Client) -> list[dict[str, Any]]:
    response = client.get(f"https://boards-api.greenhouse.io/v1/boards/{quote(slug)}/jobs", params={"content": "true"})
    response.raise_for_status()
    rows = response.json().get("jobs", [])
    return [{
        "company": company, "title": row.get("title", ""), "url": row.get("absolute_url", ""),
        "location": (row.get("location") or {}).get("name", ""), "work_mode": "Unspecified",
        "description": row.get("content", ""), "posted_at": row.get("updated_at", ""),
        "salary_text": "Not stated", "source": "Greenhouse",
    } for row in rows]


def parse_lever(slug: str, company: str, client: httpx.Client) -> list[dict[str, Any]]:
    response = client.get(f"https://api.lever.co/v0/postings/{quote(slug)}", params={"mode": "json", "limit": 100})
    response.raise_for_status()
    rows = response.json()
    output = []
    for row in rows:
        categories = row.get("categories") or {}
        salary = row.get("salaryRange") or {}
        salary_text = row.get("salaryText") or "Not stated"
        if salary and (salary.get("min") or salary.get("max")):
            salary_text = f"{salary.get('currency', 'USD')} {salary.get('min', '')} - {salary.get('max', '')} {salary.get('interval', '')}".strip()
        output.append({
            "company": company, "title": row.get("text", ""), "url": row.get("hostedUrl", ""),
            "location": categories.get("location", ""), "work_mode": row.get("workplaceType") or "Unspecified",
            "description": row.get("descriptionPlain", ""), "posted_at": str(row.get("createdAt", "")),
            "salary_text": salary_text, "source": "Lever",
        })
    return output


def parse_ashby(slug: str, company: str, client: httpx.Client) -> list[dict[str, Any]]:
    response = client.get(f"https://api.ashbyhq.com/posting-api/job-board/{quote(slug)}", params={"includeCompensation": "true"})
    response.raise_for_status()
    rows = response.json().get("jobs", [])
    output = []
    for row in rows:
        comp = row.get("compensation") or {}
        salary_text = comp.get("scrapeableCompensationSalary") or comp.get("compensationTierSummary") or "Not stated"
        output.append({
            "company": company, "title": row.get("title", ""), "url": row.get("jobUrl", ""),
            "location": row.get("location", ""), "work_mode": "Remote" if row.get("isRemote") else "Unspecified",
            "description": row.get("descriptionPlain", ""), "posted_at": row.get("publishedAt", ""),
            "salary_text": salary_text, "source": "Ashby",
        })
    return output


def sponsor_for(db: Session, company: str) -> SponsorRow | None:
    key = company_key(company)
    return db.scalar(select(SponsorRow).where(SponsorRow.company_key == key).order_by(SponsorRow.imported_at.desc()))


def score_job(job: dict[str, Any], profile: dict[str, Any], sponsor: SponsorRow | None) -> tuple[int, list[str], str, str, str]:
    title = job.get("title", "")
    text = f"{title} {job.get('description', '')}".lower()
    reasons: list[str] = []
    score = 0
    skills = [str(item).strip() for item in profile.get("skills", []) if str(item).strip()]
    skill_hits = [skill for skill in skills if skill.lower() in text]
    if skills:
        skill_score = min(42, round(42 * len(skill_hits) / max(1, min(len(skills), 8))))
        score += skill_score
        if skill_hits:
            reasons.append("Matching skills: " + ", ".join(skill_hits[:8]))
    targets = [str(item).strip() for item in profile.get("target_titles", []) if str(item).strip()]
    title_hits = [target for target in targets if target.lower() in title.lower()]
    if title_hits:
        score += 24
        reasons.append("Title matches target: " + title_hits[0])
    elif targets and any(token in title.lower() for target in targets for token in target.lower().split() if len(token) > 3):
        score += 10
        reasons.append("Title has related target-role terms")
    min_salary = profile.get("minimum_salary")
    salary_min, salary_max = parse_salary(job.get("salary_text", ""))
    if min_salary and salary_max and salary_max >= int(min_salary):
        score += 12
        reasons.append("Stated salary overlaps your minimum")
    elif min_salary and not salary_max:
        reasons.append("Salary not stated; check before applying")
    modes = [str(item).lower() for item in profile.get("work_modes", [])]
    mode_text = (job.get("work_mode", "") + " " + job.get("location", "") + " " + job.get("description", "")[:2000]).lower()
    mode_hit = next((mode for mode in modes if mode in mode_text), None)
    if mode_hit:
        score += 10
        reasons.append("Work mode signal: " + mode_hit)
    if profile.get("location_preferences"):
        loc = (job.get("location", "") + " " + job.get("description", "")[:1000]).lower()
        if any(str(place).lower() in loc for place in profile["location_preferences"] if str(place).strip()):
            score += 8
            reasons.append("Location matches a saved preference")

    low = job.get("description", "").lower()
    negative = any(term in low for term in ["unable to sponsor", "will not sponsor", "cannot sponsor", "no visa sponsorship", "not provide sponsorship"])
    explicit = any(term in low for term in ["visa sponsorship available", "will sponsor", "sponsorship is available", "sponsor h-1b", "h-1b sponsorship"])
    if negative:
        sponsor_status, sponsor_evidence, sponsor_url = "posting says no", "The posting contains a sponsorship restriction.", job.get("url", "")
    elif explicit:
        sponsor_status, sponsor_evidence, sponsor_url = "posting mentions sponsorship", "Sponsorship wording found in this job posting; verify with the employer.", job.get("url", "")
        score += 4
    elif sponsor:
        sponsor_status = "employer filing history"
        sponsor_evidence = f"{sponsor.evidence_type}; {sponsor.evidence_year or 'year not stated'}; {sponsor.filing_count} filing(s). This is employer history, not a guarantee for this job."
        sponsor_url = sponsor.evidence_url
        score += 4
    else:
        sponsor_status, sponsor_evidence, sponsor_url = "unclear", "No employer-level evidence imported and no clear sponsorship statement found in the posting.", ""
    return min(100, score), reasons, sponsor_status, sponsor_evidence, sponsor_url


def send_email(to_address: str, subject: str, body: str) -> bool:
    host = os.getenv("SMTP_HOST", "").strip()
    sender = os.getenv("SMTP_FROM", "").strip()
    if not host or not sender or not to_address:
        return False
    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = sender
    message["To"] = to_address
    message.set_content(body)
    try:
        port = int(os.getenv("SMTP_PORT", "587"))
        if port == 465:
            with smtplib.SMTP_SSL(host, port, timeout=15) as smtp:
                if os.getenv("SMTP_USERNAME"):
                    smtp.login(os.getenv("SMTP_USERNAME", ""), os.getenv("SMTP_PASSWORD", ""))
                smtp.send_message(message)
        else:
            with smtplib.SMTP(host, port, timeout=15) as smtp:
                smtp.ehlo()
                if os.getenv("SMTP_STARTTLS", "true").lower() == "true":
                    smtp.starttls()
                    smtp.ehlo()
                if os.getenv("SMTP_USERNAME"):
                    smtp.login(os.getenv("SMTP_USERNAME", ""), os.getenv("SMTP_PASSWORD", ""))
                smtp.send_message(message)
        return True
    except Exception:
        log.exception("Email notification failed")
        return False


def run_search(db: Session) -> dict[str, Any]:
    profile = ensure_profile(db).data or DEFAULT_PROFILE
    sources = db.scalars(select(SourceRow).where(SourceRow.enabled.is_(True))).all()
    result = {"sources_checked": 0, "jobs_seen": 0, "jobs_saved": 0, "new_matches": 0, "alerts_sent": 0, "source_errors": []}
    if not sources:
        result["message"] = "Add Greenhouse, Lever, or Ashby job-board sources to begin searching."
        return result
    adapters = {"greenhouse": parse_greenhouse, "lever": parse_lever, "ashby": parse_ashby}
    client = httpx.Client(timeout=20, follow_redirects=True, headers={"User-Agent": "NeuralAnt/1.0 job-search"})
    try:
        for source in sources:
            result["sources_checked"] += 1
            try:
                rows = adapters[source.ats_type](source.slug, source.company_name, client)
            except Exception as exc:
                log.warning("Job source %s/%s failed: %s", source.ats_type, source.slug, exc.__class__.__name__)
                result["source_errors"].append({"company": source.company_name, "source": source.ats_type, "error": exc.__class__.__name__})
                continue
            for raw in rows:
                url = str(raw.get("url", "")).strip()
                title = str(raw.get("title", "")).strip()
                if not url.startswith("https://") or not title:
                    continue
                result["jobs_seen"] += 1
                existing = db.scalar(select(JobRow).where(JobRow.url == url))
                sponsor = sponsor_for(db, source.company_name)
                fit, reasons, status, evidence, evidence_url = score_job(raw, profile, sponsor)
                salary_min, salary_max = parse_salary(raw.get("salary_text", ""))
                if existing:
                    existing.last_seen = datetime.now(timezone.utc)
                    existing.fit_score = fit
                    existing.fit_reasons = reasons
                    existing.sponsorship_status = status
                    existing.sponsorship_evidence = evidence
                    existing.sponsorship_url = evidence_url
                    continue
                job = JobRow(
                    url=url, company_name=source.company_name, title=title, source=str(raw.get("source", source.ats_type)),
                    location=str(raw.get("location", ""))[:500], work_mode=str(raw.get("work_mode", "Unspecified"))[:80],
                    salary_text=str(raw.get("salary_text", "Not stated"))[:255], salary_min=salary_min, salary_max=salary_max,
                    description=re.sub(r"<[^>]+>", " ", str(raw.get("description", "")))[:40000],
                    posted_at=str(raw.get("posted_at", ""))[:80], fit_score=fit, fit_reasons=reasons,
                    sponsorship_status=status, sponsorship_evidence=evidence, sponsorship_url=evidence_url,
                )
                db.add(job)
                try:
                    db.flush()
                except IntegrityError:
                    db.expunge(job)
                    continue
                result["jobs_saved"] += 1
                if fit >= int(profile.get("match_threshold") or 55):
                    result["new_matches"] += 1
                    if profile.get("email_alerts") and profile.get("email"):
                        already_sent = db.scalar(select(AlertRow).where(AlertRow.job_url == url))
                        if not already_sent:
                            body = (f"{title} — {source.company_name}\n\nFit score: {fit}%\n"
                                    f"Location: {job.location or 'Not stated'}\nWork mode: {job.work_mode}\n"
                                    f"Salary: {job.salary_text}\nSponsorship signal: {status}\n{evidence}\n\n{url}")
                            if send_email(profile["email"], f"Neural Ant match: {title} at {source.company_name}", body):
                                db.add(AlertRow(job_url=url))
                                result["alerts_sent"] += 1
        db.commit()
    finally:
        client.close()
    return result


def serialize_application(item: ApplicationRow) -> dict[str, Any]:
    return {
        "id": item.id, "stage": item.stage, "notes": item.notes, "created_at": now_iso(item.created_at),
        "job": public_job_payload(item.job),
    }


class LoginBody(BaseModel):
    password: str = Field(min_length=1, max_length=256)


class ProfileBody(BaseModel):
    model_config = {"extra": "allow"}


class SourceBody(BaseModel):
    ats_type: str
    slug: str = Field(min_length=1, max_length=160)
    company_name: str = Field(min_length=1, max_length=255)


class ApplicationBody(BaseModel):
    job_id: int
    notes: str = Field(default="", max_length=8000)


class StageBody(BaseModel):
    stage: str
    notes: str | None = Field(default=None, max_length=8000)


class TailorBody(BaseModel):
    confirmed_resume_use: bool = False
    use_external_ai: bool = False


app = FastAPI(title="Neural Ant", version="1.0.0", docs_url="/api/docs" if APP_ENV != "production" else None, redoc_url=None)


@app.middleware("http")
async def security_headers(request: Request, call_next):
    if request.method in {"POST", "PUT", "PATCH", "DELETE"} and request.url.path.startswith("/api/") and request.url.path not in {"/api/auth/login", "/api/internal/run-search"}:
        origin = request.headers.get("origin")
        if origin and origin.rstrip("/") != str(request.base_url).rstrip("/"):
            return JSONResponse({"detail": "Cross-origin requests are not allowed."}, status_code=403)
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    return response


@app.on_event("startup")
def initialize_database() -> None:
    Base.metadata.create_all(bind=engine)


@app.get("/api/health")
def health():
    return {"ok": True, "service": "neural-ant"}


@app.get("/api/config")
def config():
    return {"password_configured": bool(os.getenv("NEURAL_ANT_PASSWORD")), "ai_configured": bool(os.getenv("OPENAI_API_KEY") and os.getenv("OPENAI_MODEL")), "email_configured": bool(os.getenv("SMTP_HOST") and os.getenv("SMTP_FROM")), "resume_storage": "gcs" if GCS_BUCKET else "local volume"}


@app.get("/api/session")
def session_status(request: Request):
    if not authenticated(request):
        raise HTTPException(status_code=401, detail="Not signed in")
    return {"authenticated": True}


_login_attempts: dict[str, list[float]] = {}


@app.post("/api/auth/login")
def login(body: LoginBody, request: Request):
    expected = os.getenv("NEURAL_ANT_PASSWORD", "")
    if not expected:
        raise HTTPException(status_code=503, detail="Set NEURAL_ANT_PASSWORD in the server environment first.")
    client_ip = request.client.host if request.client else "unknown"
    import time

    attempts = [stamp for stamp in _login_attempts.get(client_ip, []) if time.time() - stamp < 300]
    if len(attempts) >= 10:
        raise HTTPException(status_code=429, detail="Too many sign-in attempts. Wait a few minutes and retry.")
    attempts.append(time.time())
    _login_attempts[client_ip] = attempts
    if not hmac.compare_digest(body.password.encode("utf-8"), expected.encode("utf-8")):
        raise HTTPException(status_code=401, detail="Password is incorrect.")
    response = JSONResponse({"authenticated": True})
    response.set_cookie(COOKIE_NAME, serializer.dumps({"sub": "owner"}), httponly=True, secure=APP_ENV == "production", samesite="strict", max_age=SESSION_MAX_AGE, path="/")
    _login_attempts.pop(client_ip, None)
    return response


@app.post("/api/auth/logout")
def logout():
    response = JSONResponse({"authenticated": False})
    response.delete_cookie(COOKIE_NAME, path="/", httponly=True, secure=APP_ENV == "production", samesite="strict")
    return response


@app.get("/api/profile")
def get_profile(_: None = Depends(require_auth), db: Session = Depends(get_db)):
    row = ensure_profile(db)
    resume = db.get(ResumeRow, 1)
    return {"profile": row.data or DEFAULT_PROFILE, "resume": {"filename": resume.filename, "uploaded_at": now_iso(resume.uploaded_at)} if resume else None}


@app.put("/api/profile")
def save_profile(body: ProfileBody, _: None = Depends(require_auth), db: Session = Depends(get_db)):
    allowed = set(DEFAULT_PROFILE)
    values = {key: value for key, value in body.model_dump().items() if key in allowed}
    if "email" in values and values["email"] and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", str(values["email"])):
        raise HTTPException(status_code=422, detail="Enter a valid notification email address.")
    if "minimum_salary" in values and values["minimum_salary"] not in (None, ""):
        values["minimum_salary"] = int(values["minimum_salary"])
    for key in ("skills", "target_titles", "location_preferences", "work_modes"):
        if key in values and not isinstance(values[key], list):
            values[key] = [part.strip() for part in str(values[key]).split(",") if part.strip()]
    if "match_threshold" in values:
        values["match_threshold"] = max(0, min(100, int(values["match_threshold"])))
    row = ensure_profile(db)
    row.data = {**DEFAULT_PROFILE, **(row.data or {}), **values}
    row.updated_at = datetime.now(timezone.utc)
    db.commit()
    return {"profile": row.data}


@app.post("/api/resume")
async def upload_resume(file: UploadFile = File(...), _: None = Depends(require_auth), db: Session = Depends(get_db)):
    filename = Path(file.filename or "resume").name
    if Path(filename).suffix.lower() not in {".pdf", ".docx"}:
        raise HTTPException(status_code=415, detail="Upload a PDF or DOCX resume.")
    content = await file.read(MAX_RESUME_BYTES + 1)
    if len(content) > MAX_RESUME_BYTES:
        raise HTTPException(status_code=413, detail="Resume must be 10 MB or smaller.")
    text = extract_resume_text(filename, content)
    if len(text) < 80:
        raise HTTPException(status_code=422, detail="Could not extract enough text. Try a text-based PDF or DOCX file.")
    file_key = uuid.uuid4().hex
    file_ref = put_resume_file(file_key, filename, content)
    existing = db.get(ResumeRow, 1)
    if existing:
        delete_resume_file(existing.file_key)
        existing.filename = filename
        existing.file_key = file_ref
        existing.file_sha256 = hashlib.sha256(content).hexdigest()
        existing.extracted_text = text
        existing.uploaded_at = datetime.now(timezone.utc)
    else:
        db.add(ResumeRow(id=1, filename=filename, file_key=file_ref, file_sha256=hashlib.sha256(content).hexdigest(), extracted_text=text))
    extracted = extract_profile(text)
    profile_row = ensure_profile(db)
    profile_data = {**DEFAULT_PROFILE, **(profile_row.data or {})}
    for key in ("name", "email", "skills"):
        if extracted.get(key) and not profile_data.get(key):
            profile_data[key] = extracted[key]
    profile_row.data = profile_data
    db.commit()
    return {"filename": filename, "characters_extracted": len(text), "extracted_fields": extracted, "message": "Resume stored privately. Review extracted name, email, and skills in My profile before using them."}


@app.delete("/api/resume")
def delete_resume(_: None = Depends(require_auth), db: Session = Depends(get_db)):
    row = db.get(ResumeRow, 1)
    if row:
        delete_resume_file(row.file_key)
        db.delete(row)
        db.commit()
    return {"deleted": True}


@app.get("/api/sources")
def list_sources(_: None = Depends(require_auth), db: Session = Depends(get_db)):
    rows = db.scalars(select(SourceRow).order_by(SourceRow.company_name)).all()
    return [{"id": row.id, "ats_type": row.ats_type, "slug": row.slug, "company_name": row.company_name, "enabled": row.enabled} for row in rows]


@app.post("/api/sources")
def add_source(body: SourceBody, _: None = Depends(require_auth), db: Session = Depends(get_db)):
    ats_type = body.ats_type.lower().strip()
    slug = body.slug.strip()
    company = body.company_name.strip()
    if ats_type not in SUPPORTED_ATS:
        raise HTTPException(status_code=422, detail="Choose Greenhouse, Lever, or Ashby.")
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,159}", slug):
        raise HTTPException(status_code=422, detail="Board slug contains unsupported characters.")
    row = SourceRow(ats_type=ats_type, slug=slug, company_name=company)
    db.add(row)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="That job board is already configured.") from exc
    db.refresh(row)
    return {"id": row.id, "ats_type": row.ats_type, "slug": row.slug, "company_name": row.company_name, "enabled": row.enabled}


@app.delete("/api/sources/{source_id}")
def remove_source(source_id: int, _: None = Depends(require_auth), db: Session = Depends(get_db)):
    row = db.get(SourceRow, source_id)
    if not row:
        raise HTTPException(status_code=404, detail="Source not found.")
    db.delete(row)
    db.commit()
    return {"deleted": True}


@app.post("/api/sponsors/import")
async def import_sponsors(file: UploadFile = File(...), _: None = Depends(require_auth), db: Session = Depends(get_db)):
    if not (file.filename or "").lower().endswith(".csv"):
        raise HTTPException(status_code=415, detail="Upload a normalized CSV file.")
    raw = await file.read(5 * 1024 * 1024 + 1)
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="CSV must be 5 MB or smaller.")
    try:
        reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))
    except UnicodeDecodeError as exc:
        raise HTTPException(status_code=400, detail="CSV must be UTF-8 encoded.") from exc
    required = {"company_name", "evidence_url", "evidence_year", "filing_count"}
    if not reader.fieldnames or not required.issubset(set(reader.fieldnames)):
        raise HTTPException(status_code=422, detail="CSV headers must include company_name,evidence_url,evidence_year,filing_count. Optional: evidence_type,notes.")
    imported = 0
    for item in reader:
        name = (item.get("company_name") or "").strip()
        url = (item.get("evidence_url") or "").strip()
        if not name or not url.startswith("https://"):
            continue
        key = company_key(name)
        existing = db.scalar(select(SponsorRow).where(SponsorRow.company_key == key, SponsorRow.evidence_year == (item.get("evidence_year") or "").strip(), SponsorRow.evidence_url == url))
        if existing:
            continue
        try:
            count = max(0, int(item.get("filing_count") or 0))
        except ValueError:
            count = 0
        db.add(SponsorRow(company_key=key, company_name=name, evidence_type=(item.get("evidence_type") or "public filing")[:80], evidence_url=url, evidence_year=(item.get("evidence_year") or "")[:16], filing_count=count, notes=(item.get("notes") or "")[:2000]))
        imported += 1
    db.commit()
    return {"imported": imported, "message": "Employer filing history is only a historical signal; it does not guarantee sponsorship for a specific role."}


@app.get("/api/sponsors")
def list_sponsors(_: None = Depends(require_auth), db: Session = Depends(get_db)):
    rows = db.scalars(select(SponsorRow).order_by(SponsorRow.company_name)).all()
    return [{"company_name": row.company_name, "evidence_type": row.evidence_type, "evidence_url": row.evidence_url, "evidence_year": row.evidence_year, "filing_count": row.filing_count, "notes": row.notes} for row in rows]


@app.post("/api/search/run")
def search_now(_: None = Depends(require_auth), db: Session = Depends(get_db)):
    return run_search(db)


@app.post("/api/internal/run-search")
def scheduled_search(x_cron_secret: str | None = Header(default=None), db: Session = Depends(get_db)):
    expected = os.getenv("CRON_SECRET", "")
    if not expected:
        raise HTTPException(status_code=503, detail="CRON_SECRET is not configured.")
    if not x_cron_secret or not hmac.compare_digest(x_cron_secret, expected):
        raise HTTPException(status_code=403, detail="Invalid scheduled-job credential.")
    return run_search(db)


@app.get("/api/jobs")
def get_jobs(min_score: int = 0, _: None = Depends(require_auth), db: Session = Depends(get_db)):
    rows = db.scalars(select(JobRow).where(JobRow.fit_score >= max(0, min_score)).order_by(JobRow.fit_score.desc(), JobRow.first_seen.desc()).limit(300)).all()
    return [public_job_payload(row) for row in rows]


@app.post("/api/applications")
def create_application(body: ApplicationBody, _: None = Depends(require_auth), db: Session = Depends(get_db)):
    job = db.get(JobRow, body.job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found.")
    existing = db.scalar(select(ApplicationRow).where(ApplicationRow.job_id == body.job_id))
    if existing:
        return serialize_application(existing)
    row = ApplicationRow(job_id=job.id, notes=body.notes.strip())
    db.add(row)
    db.commit()
    db.refresh(row)
    return serialize_application(row)


@app.get("/api/applications")
def list_applications(_: None = Depends(require_auth), db: Session = Depends(get_db)):
    rows = db.scalars(select(ApplicationRow).order_by(ApplicationRow.updated_at.desc())).all()
    return [serialize_application(row) for row in rows]


@app.patch("/api/applications/{application_id}")
def update_application(application_id: int, body: StageBody, _: None = Depends(require_auth), db: Session = Depends(get_db)):
    if body.stage not in STAGES:
        raise HTTPException(status_code=422, detail="Choose a supported application stage.")
    row = db.get(ApplicationRow, application_id)
    if not row:
        raise HTTPException(status_code=404, detail="Application not found.")
    row.stage = body.stage
    if body.notes is not None:
        row.notes = body.notes.strip()
    row.updated_at = datetime.now(timezone.utc)
    db.commit()
    return serialize_application(row)


def deterministic_tailor(resume_text: str, job: JobRow, profile: dict[str, Any]) -> str:
    skills = [str(value) for value in profile.get("skills", [])]
    terms = list(dict.fromkeys(skills + [word for word in re.findall(r"[A-Za-z][A-Za-z+#.-]{2,}", job.title + " " + job.description[:4000]) if len(word) > 3]))
    lines = [line.strip() for line in resume_text.splitlines() if len(line.strip()) > 18]
    selected = [line for line in lines if any(term.lower() in line.lower() for term in terms[:80])][:18]
    if not selected:
        selected = lines[:12]
    return (
        f"# Resume draft for {job.title} — {job.company_name}\n\n"
        "Review every line against your source resume before using this draft. This local draft only reorders existing lines; it does not add claims.\n\n"
        "## Relevant existing resume content\n\n" + "\n".join(f"- {line}" for line in selected) +
        "\n\n## Role details to verify\n\n"
        f"- Posting: {job.url}\n- Compensation: {job.salary_text}\n- Sponsorship signal: {job.sponsorship_status} — {job.sponsorship_evidence}\n"
    )


def external_tailor(resume_text: str, job: JobRow, profile: dict[str, Any]) -> tuple[str, str]:
    api_key = os.getenv("OPENAI_API_KEY", "")
    model = os.getenv("OPENAI_MODEL", "")
    if not api_key or not model:
        raise HTTPException(status_code=503, detail="External AI is not configured. Set OPENAI_API_KEY and OPENAI_MODEL, or use the local draft option.")
    base_url = os.getenv("OPENAI_BASE_URL", "https://api.openai.com/v1").rstrip("/")
    prompt = f"""Create an ATS-friendly resume draft tailored to this job. Only use facts present in the source resume. Never invent employers, dates, skills, metrics, degrees, or work authorization. Preserve the candidate's actual meaning. Mark any suggested gap or missing detail as a question instead of filling it in. Return Markdown and include a short 'Facts to verify' section.

PROFILE:
{json.dumps(profile, ensure_ascii=False)[:8000]}

JOB POSTING:
Title: {job.title}
Company: {job.company_name}
Description: {job.description[:12000]}

SOURCE RESUME:
{resume_text[:20000]}
"""
    try:
        response = httpx.post(f"{base_url}/chat/completions", headers={"Authorization": f"Bearer {api_key}"}, json={"model": model, "temperature": 0.2, "messages": [{"role": "system", "content": "You are a careful resume editor. Never fabricate candidate facts."}, {"role": "user", "content": prompt}]}, timeout=60)
        response.raise_for_status()
        content = response.json()["choices"][0]["message"]["content"]
        return str(content), model
    except httpx.HTTPError as exc:
        log.exception("Configured AI provider call failed")
        raise HTTPException(status_code=502, detail="The configured AI provider could not complete this draft.") from exc
    except (KeyError, IndexError, TypeError) as exc:
        raise HTTPException(status_code=502, detail="The configured AI provider returned an unexpected response.") from exc


@app.post("/api/jobs/{job_id}/tailor")
def tailor_resume(job_id: int, body: TailorBody, _: None = Depends(require_auth), db: Session = Depends(get_db)):
    if not body.confirmed_resume_use:
        raise HTTPException(status_code=400, detail="Confirm that Neural Ant may use your stored resume to prepare this draft.")
    resume = db.get(ResumeRow, 1)
    if not resume:
        raise HTTPException(status_code=409, detail="Upload a resume first.")
    job = db.get(JobRow, job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found.")
    profile = ensure_profile(db).data or DEFAULT_PROFILE
    if body.use_external_ai:
        content, model = external_tailor(resume.extracted_text, job, profile)
    else:
        content, model = deterministic_tailor(resume.extracted_text, job, profile), "local-match-draft"
    draft = DraftRow(job_id=job.id, content=content, model=model)
    db.add(draft)
    db.commit()
    db.refresh(draft)
    return {"id": draft.id, "content": draft.content, "model": draft.model, "created_at": now_iso(draft.created_at), "warning": "Draft only. Review every claim and submit applications yourself."}


@app.get("/api/drafts")
def list_drafts(_: None = Depends(require_auth), db: Session = Depends(get_db)):
    rows = db.scalars(select(DraftRow).order_by(DraftRow.created_at.desc()).limit(100)).all()
    return [{"id": row.id, "job_id": row.job_id, "content": row.content, "model": row.model, "created_at": now_iso(row.created_at)} for row in rows]


@app.get("/api/status")
def status(_: None = Depends(require_auth), db: Session = Depends(get_db)):
    return {
        "database": "connected",
        "resume_uploaded": db.get(ResumeRow, 1) is not None,
        "job_sources": db.scalar(select(func.count(SourceRow.id))) or 0,
        "sponsor_records": db.scalar(select(func.count(SponsorRow.id))) or 0,
        "ai_configured": bool(os.getenv("OPENAI_API_KEY") and os.getenv("OPENAI_MODEL")),
        "email_configured": bool(os.getenv("SMTP_HOST") and os.getenv("SMTP_FROM")),
        "scheduled_search_configured": bool(os.getenv("CRON_SECRET")),
        "resume_storage": "private Google Cloud Storage bucket" if GCS_BUCKET else "private persistent volume",
    }


@app.get("/", include_in_schema=False)
def home():
    index = DIST_DIR / "index.html"
    if not index.exists():
        raise HTTPException(status_code=503, detail="Build the frontend with npm run build first.")
    return FileResponse(index)


if DIST_DIR.exists():
    app.mount("/", StaticFiles(directory=DIST_DIR, html=True), name="web")
